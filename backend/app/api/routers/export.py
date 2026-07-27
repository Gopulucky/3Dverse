from typing import Any
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from sqlalchemy.orm import Session
from io import BytesIO

from app.api import deps
from app.models.user import User, UserRole
from app.models.product import Product
from app.services import metrics, filament

router = APIRouter()


@router.get("/excel")
def export_excel(
    db: Session = Depends(deps.get_db),
    current_user: User = Depends(deps.get_current_active_user),
) -> Any:
    """Export creator metrics and filament status as a styled Excel workbook."""
    if current_user.role != UserRole.creator:
        raise HTTPException(status_code=403, detail="Only creators can export their data")

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()

    # --- Sheet 1: Revenue Summary ---
    ws = wb.active
    ws.title = "Revenue Summary"

    header_font = Font(bold=True, size=12, color="FFFFFF")
    header_fill = PatternFill(start_color="6E5BFF", end_color="6E5BFF", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    summary = metrics.creator_financial_summary(db, creator_id=current_user.id)

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 20

    ws.append(["Metric", "Value"])
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    metric_labels = {
        "total_sales": "Total Sales (₹)",
        "total_earnings": "Total Earnings (₹)",
        "orders_count": "Total Orders",
        "average_order_value": "Average Order Value (₹)",
        "platform_fees_paid": "Platform Fees Paid (₹)",
    }

    for key, label in metric_labels.items():
        value = summary.get(key, 0)
        if key != "orders_count":
            value = f"₹{value:,.2f}" if isinstance(value, (int, float)) else value
        row = ws.append([label, value])

    for row in ws.iter_rows(min_row=2, max_col=2):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="left")

    # --- Sheet 2: Filament Status ---
    ws2 = wb.create_sheet("Filament Status")
    ws2.column_dimensions["A"].width = 15
    ws2.column_dimensions["B"].width = 20
    ws2.column_dimensions["C"].width = 15

    ws2.append(["Product ID", "Product Title", "Whole Units Remaining"])
    for cell in ws2[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    products = db.query(Product).filter_by(creator_id=current_user.id).all()
    for product in products:
        units = filament.whole_units_remaining(db, product.id)
        ws2.append([product.id, product.title, units])

    for row in ws2.iter_rows(min_row=2, max_col=3):
        for cell in row:
            cell.border = thin_border

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=3dverse-metrics-{current_user.id}.xlsx"},
    )
